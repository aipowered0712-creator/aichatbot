"""知汇 AI 的统一模型路由 API。

用户在网页选择 local 或 online 后，前端统一调用 POST /api/chat；
此服务负责复用同一个本地 FAISS 检索器，并路由至相应的生成模型。
模型均为懒加载，不会在启动时同时占用资源。
"""

from __future__ import annotations

import asyncio
import faulthandler
import hashlib
import json
import os
import re
import time
import shutil
import tempfile
import threading
import urllib.error
import urllib.request
from contextlib import asynccontextmanager
from functools import lru_cache
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from langchain_core.embeddings import Embeddings
from pydantic import BaseModel, Field
from zhihui_chunking import STATUS


APP_DIR = Path(__file__).resolve().parent
LOCAL_MODEL_NAME = "Kimi-VL-A3B-Instruct"
ONLINE_MODEL_NAME = os.environ.get("HKPC_LLM_MODEL", "public/qwen3.8-27b")
ONLINE_BASE_URL = os.environ.get("HKPC_LLM_BASE_URL", "https://api-davinci.hkpc.org/v1/")
ONLINE_MODELS = {
    "public/qwen3.8-27b": {"label": "Qwen3.8-27B", "max_tokens": 512},
    "public/deepseek-v4-flash-w8a8-mtp": {"label": "DeepSeek-V4-Flash-W8A8-MTP", "max_tokens": 512},
    "public/minimax-m3": {"label": "MiniMax-M3", "max_tokens": 512},
    "public/qwen3.8-max-0902": {"label": "Qwen3.8-Max-0902", "max_tokens": 512, "commercial": True},
}

LOCAL_MODEL_PATH = os.environ.get("LOCAL_MODEL_PATH", r"D:\A_model\Kimi-VL-A3B-Instruct")
LIBRARY_ID_RE = re.compile(r"^[a-z0-9][a-z0-9_-]{0,40}$")
BUILTIN_LIBRARIES = {"moldpdf": "注塑設計", "metal": "金屬設計"}
ALL_SCOPES = {"全部知识空间", "全部知識空間", "全部文档", "全部文檔", "all", "*"}
NONE_SCOPES = {"不使用知识库", "不使用知識庫", "none"}
ONLINE_API_KEY = os.environ.get("HKPC_LLM_API_KEY", "")
LOCAL_EMBEDDING_MODEL = "Qwen/Qwen3-Embedding-0.6B"
LOCAL_RERANKER_MODEL = "Qwen/Qwen3-Reranker-8B"
ONLINE_EMBEDDING_MODEL = "public/qwen3-embedding-0.6b"
ONLINE_RERANKER_MODEL = "public/qwen3-reranker-8b"
EMBEDDING_BATCH_SIZE = max(1, int(os.environ.get("HKPC_EMBEDDING_BATCH", "32")))


def _env_int(name: str, default: int) -> int:
    raw = os.environ.get(name, "").strip()
    if not raw:
        return default
    try:
        return int(raw)
    except ValueError:
        return default


def clamp_chunk_size(value: int) -> int:
    return max(200, min(4000, int(value)))


def clamp_chunk_overlap(value: int, chunk_size: int) -> int:
    return max(0, min(int(value), max(0, int(chunk_size) // 2)))


def clamp_top_k(value: int) -> int:
    return max(1, min(20, int(value)))


RAG_CHUNK_SIZE = clamp_chunk_size(_env_int("HKPC_CHUNK_SIZE", 800))
RAG_CHUNK_OVERLAP = clamp_chunk_overlap(_env_int("HKPC_CHUNK_OVERLAP", 120), RAG_CHUNK_SIZE)
RAG_TOP_K = clamp_top_k(_env_int("HKPC_TOP_K", 4))


def chunking_token() -> str:
    if RAG_CHUNK_SIZE == 800 and RAG_CHUNK_OVERLAP == 120:
        return "chunking:zhihui-v1"
    return f"chunking:zhihui-v1:{RAG_CHUNK_SIZE}:{RAG_CHUNK_OVERLAP}"


def use_online_vectors() -> bool:
    override = os.environ.get("HKPC_VECTOR_BACKEND", "").strip().lower()
    if override in {"online", "hkpc", "cloud"}:
        return True
    if override in {"local", "cpu", "gpu"}:
        return False
    return bool(ONLINE_API_KEY)


def active_embedding_model() -> str:
    name = os.environ.get("HKPC_EMBEDDING_MODEL", "").strip()
    if use_online_vectors():
        if not name or name.startswith("Qwen/"):
            return ONLINE_EMBEDDING_MODEL
        return name
    return name or LOCAL_EMBEDDING_MODEL


def active_reranker_model() -> str:
    name = os.environ.get("HKPC_RERANKER_MODEL", "").strip()
    if use_online_vectors():
        if not name or name.startswith("Qwen/"):
            return ONLINE_RERANKER_MODEL
        return name
    return name or LOCAL_RERANKER_MODEL


EMBEDDING_MODEL_NAME = active_embedding_model()
RERANKER_MODEL_NAME = active_reranker_model()


INDEX_FILES = ("index.faiss", "index.pkl")


def path_is_ascii(path: Path) -> bool:
    try:
        str(path).encode("ascii")
        return True
    except UnicodeEncodeError:
        return False


def resolve_index_dir() -> Path:
    """索引文件默认放在 RAG 项目目录下的 faiss/。可用 FAISS_INDEX_DIR 覆盖。"""
    env_dir = os.environ.get("FAISS_INDEX_DIR", "").strip()
    if env_dir:
        return Path(env_dir)
    return APP_DIR / "faiss"


def same_path(left: Path, right: Path) -> bool:
    try:
        return left.resolve() == right.resolve()
    except OSError:
        return os.path.normcase(str(left)) == os.path.normcase(str(right))


def faiss_native_dir(library_id: str) -> Path:
    """FAISS C++ 在 Windows 上不能直接讀寫中文路徑，非 ASCII 目錄經 ASCII 中轉。"""
    target = library_index_dir(library_id)
    if path_is_ascii(target):
        return target
    return Path(tempfile.gettempdir()) / "zhihui-ai-faiss" / safe_library_id(library_id)


def copy_index_files(source: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    for name in INDEX_FILES:
        src = source / name
        if src.is_file():
            shutil.copy2(src, destination / name)


INDEX_DIR = resolve_index_dir()
SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm", ".pptx",
    ".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff",
}
SOURCE_TITLES = {
    "mold.pdf": "塑料注塑技術手冊",
    "mold.docx": "塑料注塑技術手冊",
    "mold-2.pdf": "塑料注塑技術手冊（續）",
    "en-ebook_-injection-molding-design-guide.pdf": "Injection Molding Design Guide (Xometry)",
    "d810tvo119rxmaubzwdb_dyneon_injection_moulding_handbook.pdf": "3M Dyneon Fluoroplastics Injection Moulding Handbook",
    "preview-9781569908167_a42563111.pdf": "Injection Mold Design Handbook (Catoen / Rees)",
    "iso-20457-2018.pdf": "ISO 20457:2018 塑料模製件公差與驗收條件",
    "injection_molding_design_references.md": "塑料注塑設計規範與重要參考書目",
}
HEADER_LINE = re.compile(
    r"^(塑料注塑技[術术]手[冊册]|Injection Moulding.*Handbook|删除\s*\[.*)$",
    re.IGNORECASE,
)

_index_lock = threading.Lock()
_stores: dict = {}
_fingerprints: dict[str, str] = {}
_statuses: dict[str, str] = {}
_errors: dict[str, str] = {}
_generations: dict[str, int] = {}
_active_library = ""
_local_llm_lock = threading.Lock()
_local_llm = None
_local_llm_status = "idle"
_local_llm_error = ""

SYSTEM_PROMPT = """你是嚴謹的企業知識庫助理。請一律使用繁體中文回答，不要使用簡體字。請嚴格基於【參考資料】回答問題。
規則：
1. 不得編造參考資料中沒有的資訊。
2. 資料不足時，明確回覆無法從現有文檔確認，並指出需要補充什麼。
3. 回答結構清晰，重要數據和條件應保留原意。
4. 若問題指定了知識庫，只依據該知識庫的資料，不要混入其他知識庫。
5. 排版要能直接閱讀：每個小節單獨用一行「## 標題」。條目用「- 」開頭，只寫一層，不要巢狀，不要用星號。關鍵溫度與數值寫成 **260℃** 這種粗體。不要用三級標題、表格或代碼塊。

【參考資料】
{context}

【用戶問題】
{question}

回答："""

@asynccontextmanager
async def lifespan(_app: FastAPI):
    faulthandler.enable()
    ensure_builtin_libraries()
    threading.Thread(target=warmup_index, name="zhihui-index", daemon=True).start()
    yield


app = FastAPI(title="知匯 AI API", version="1.0.0", lifespan=lifespan)


@app.middleware("http")
async def disable_frontend_cache(request, call_next):
    response = await call_next(request)
    if request.url.path in {"/", "/index.html", "/app.js", "/styles.css"}:
        response.headers["Cache-Control"] = "no-store, max-age=0"
    return response


class ChatRequest(BaseModel):
    question: str = Field(min_length=1, max_length=8000)
    provider: Literal["local", "online"] = "local"
    model: str | None = None
    knowledge_scope: str = "全部知識空間"
    use_knowledge: bool = True
    chunk_size: int | None = Field(default=None, ge=200, le=4000)
    chunk_overlap: int | None = Field(default=None, ge=0, le=2000)
    top_k: int | None = Field(default=None, ge=1, le=20)
    custom_prompt: str | None = Field(default=None, max_length=8000)
    debug: bool = False


def path_exists(path: str) -> bool:
    return bool(path) and Path(path).exists()


def knowledge_root() -> Path:
    configured = Path(os.environ.get("KNOWLEDGE_BASE_DIR", str(APP_DIR / "knowledge_base" / "moldpdf")))
    if configured.name == "knowledge_base":
        return configured
    if configured.parent.name == "knowledge_base":
        return configured.parent
    return APP_DIR / "knowledge_base"


def safe_library_id(value: str) -> str:
    cleaned = re.sub(r"[^a-z0-9_-]+", "-", (value or "").strip().lower()).strip("-")
    if not cleaned or not LIBRARY_ID_RE.match(cleaned):
        raise ValueError("知識庫編號只可使用英數、連字號或底線")
    return cleaned


def library_dir(library_id: str) -> Path:
    return knowledge_root() / safe_library_id(library_id)


def library_index_dir(library_id: str) -> Path:
    return INDEX_DIR / safe_library_id(library_id)


def read_library_name(library_id: str) -> str:
    fallback = BUILTIN_LIBRARIES.get(library_id, library_id)
    meta_path = library_dir(library_id) / "library.json"
    if not meta_path.is_file():
        return fallback
    try:
        data = json.loads(meta_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return fallback
    return str(data.get("name") or "").strip() or fallback


def write_library_meta(library_id: str, name: str) -> None:
    folder = library_dir(library_id)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "library.json").write_text(
        json.dumps({"id": library_id, "name": name}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def slug_from_name(name: str) -> str:
    ascii_part = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    if ascii_part and LIBRARY_ID_RE.match(ascii_part):
        return ascii_part[:40]
    return "kb-" + hashlib.sha1(name.strip().encode("utf-8")).hexdigest()[:8]


def ensure_builtin_libraries() -> None:
    root = knowledge_root()
    root.mkdir(parents=True, exist_ok=True)
    existing = [
        path for path in root.iterdir()
        if path.is_dir() and not path.name.startswith(".") and LIBRARY_ID_RE.match(path.name)
    ]
    if existing:
        return
    for library_id, name in BUILTIN_LIBRARIES.items():
        write_library_meta(library_id, name)
    metal_readme = library_dir("metal") / "README.md"
    if not metal_readme.is_file():
        metal_readme.write_text(
            "# 金屬設計\n\n此知識庫與注塑設計分開建立索引。把金屬設計文件上傳到這裡即可檢索。\n",
            encoding="utf-8",
        )


def list_library_ids() -> list[str]:
    root = knowledge_root()
    if not root.is_dir():
        return []
    return sorted(
        path.name
        for path in root.iterdir()
        if path.is_dir() and not path.name.startswith(".") and LIBRARY_ID_RE.match(path.name)
    )


def iter_library_files(library_id: str) -> list[Path]:
    root = library_dir(library_id)
    if not root.is_dir():
        return []
    return sorted(
        path for path in root.rglob("*")
        if (
            path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
            and path.name.lower() not in {"readme.md", "library.json"}
            and not any(part.startswith(".") for part in path.relative_to(root).parts)
        )
    )


def knowledge_files(library_id: str | None = None):
    if library_id:
        return iter_library_files(library_id)
    files = []
    for item in list_library_ids():
        files.extend(iter_library_files(item))
    return files


def libraries_with_files() -> list[str]:
    return [item for item in list_library_ids() if iter_library_files(item)]


def file_fingerprint(library_id: str) -> str:
    files = iter_library_files(library_id)
    if not files:
        return ""
    root = library_dir(library_id)
    parts = [
        f"library:{library_id}",
        f"embedding-model:{active_embedding_model()}",
        f"vector-backend:{'online' if use_online_vectors() else 'local'}",
        chunking_token(),
    ]
    parts.extend(
        f"{path.relative_to(root).as_posix()}:{path.stat().st_mtime_ns}:{path.stat().st_size}"
        for path in files
    )
    return hashlib.sha1("\n".join(parts).encode("utf-8")).hexdigest()


def index_is_current(library_id: str) -> bool:
    folder = library_index_dir(library_id)
    marker = folder / "fingerprint.txt"
    fingerprint = file_fingerprint(library_id)
    return bool(
        fingerprint
        and marker.is_file()
        and (folder / "index.faiss").is_file()
        and marker.read_text(encoding="utf-8").strip() == fingerprint
    )


def resolve_library_id(scope: str) -> str:
    wanted = (scope or "").strip()
    for library_id in list_library_ids():
        if wanted in {library_id, read_library_name(library_id)}:
            return library_id
    raise RuntimeError(f"找不到知識庫：{wanted}")


def clean_text(text: str) -> str:
    lines = [line for line in text.splitlines() if line.strip() and not HEADER_LINE.match(line.strip())]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def source_title(source_path: str) -> str:
    name = Path(source_path).name
    return SOURCE_TITLES.get(name.lower(), name) or "知識庫文檔"


@lru_cache(maxsize=1)
def resolve_torch_device() -> str:
    override = os.environ.get("HKPC_TORCH_DEVICE", "").strip()
    if override:
        return override
    try:
        import torch
    except ImportError:
        return "cpu"
    if not torch.cuda.is_available():
        return "cpu"
    best_index, best_free = 0, -1
    for index in range(torch.cuda.device_count()):
        try:
            free, _total = torch.cuda.mem_get_info(index)
        except Exception:
            free = 0
        if free > best_free:
            best_free, best_index = free, index
    return f"cuda:{best_index}"


def describe_torch_device(device: str) -> str:
    if not device.startswith("cuda"):
        return device
    try:
        import torch

        index = int(device.split(":")[1]) if ":" in device else torch.cuda.current_device()
        return f"{device}（{torch.cuda.get_device_name(index)}）"
    except Exception:
        return device


def update_env_file(updates: dict[str, str]) -> None:
    path = APP_DIR / ".env"
    lines = path.read_text(encoding="utf-8").splitlines() if path.is_file() else []
    seen: set[str] = set()
    rewritten: list[str] = []
    for line in lines:
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in line:
            rewritten.append(line)
            continue
        key = line.split("=", 1)[0].strip()
        if key.startswith("export "):
            key = key[7:].strip()
        if key in updates:
            rewritten.append(f"{key}={updates[key]}")
            seen.add(key)
        else:
            rewritten.append(line)
    for key, value in updates.items():
        if key not in seen:
            rewritten.append(f"{key}={value}")
    path.write_text("\n".join(rewritten) + "\n", encoding="utf-8")


MODEL_CATALOG_PATH = APP_DIR / "online_models.json"
MODEL_ID_RE = re.compile(r"^[A-Za-z0-9_./:-]{1,120}$")


def public_model_rows() -> list[dict]:
    return [
        {
            "id": model_id,
            "label": spec["label"],
            "commercial": bool(spec.get("commercial")),
            "api_key_set": bool(spec.get("api_key")),
            "base_url": str(spec.get("base_url") or ""),
        }
        for model_id, spec in ONLINE_MODELS.items()
    ]


def model_base_url(model_id: str) -> str:
    spec = ONLINE_MODELS.get(model_id) or {}
    return str(spec.get("base_url") or ONLINE_BASE_URL).rstrip("/")


def model_api_key(model_id: str) -> str:
    spec = ONLINE_MODELS.get(model_id) or {}
    return str(spec.get("api_key") or ONLINE_API_KEY or "")


def load_model_catalog() -> None:
    if not MODEL_CATALOG_PATH.is_file():
        return
    try:
        data = json.loads(MODEL_CATALOG_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return
    loaded: dict[str, dict] = {}
    for item in data.get("models") or []:
        model_id = str(item.get("id") or "").strip()
        label = str(item.get("label") or "").strip() or model_id
        if MODEL_ID_RE.fullmatch(model_id):
            loaded[model_id] = {
                "label": label[:40],
                "max_tokens": 512,
                "commercial": bool(item.get("commercial")),
                "api_key": str(item.get("api_key") or "").strip(),
                "base_url": str(item.get("base_url") or "").strip().rstrip("/"),
            }
    if loaded:
        ONLINE_MODELS.clear()
        ONLINE_MODELS.update(loaded)


def save_model_catalog() -> None:
    models = []
    for model_id, spec in ONLINE_MODELS.items():
        item = {
            "id": model_id,
            "label": spec["label"],
            "commercial": bool(spec.get("commercial")),
        }
        if spec.get("api_key"):
            item["api_key"] = spec["api_key"]
        if spec.get("base_url"):
            item["base_url"] = spec["base_url"]
        models.append(item)
    MODEL_CATALOG_PATH.write_text(
        json.dumps({"models": models}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


load_model_catalog()


def apply_online_settings(api_key: str | None = None, base_url: str | None = None, model: str | None = None) -> None:
    global ONLINE_API_KEY, ONLINE_BASE_URL, ONLINE_MODEL_NAME
    updates: dict[str, str] = {}
    if api_key:
        ONLINE_API_KEY = api_key
        os.environ["HKPC_LLM_API_KEY"] = api_key
        updates["HKPC_LLM_API_KEY"] = api_key
    if base_url:
        ONLINE_BASE_URL = base_url
        os.environ["HKPC_LLM_BASE_URL"] = base_url
        updates["HKPC_LLM_BASE_URL"] = base_url
    if model:
        if model not in ONLINE_MODELS:
            ONLINE_MODELS[model] = {"label": model, "max_tokens": 512}
        ONLINE_MODEL_NAME = model
        os.environ["HKPC_LLM_MODEL"] = model
        updates["HKPC_LLM_MODEL"] = model
    if updates:
        update_env_file(updates)


def hkpc_post(path: str, payload: dict, timeout: int = 120, api_key: str | None = None, base_url: str | None = None) -> dict:
    token = api_key or ONLINE_API_KEY
    root = (base_url or ONLINE_BASE_URL).rstrip("/")
    if not token:
        raise RuntimeError("尚未設定此模型的 API Key")
    if not root:
        raise RuntimeError("尚未設定此模型的 URL")
    request = urllib.request.Request(
        root + path,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HKPC 接口调用失败：{error.code} {detail[:300]}") from error


def _l2_normalize(vector: list[float]) -> list[float]:
    norm = sum(value * value for value in vector) ** 0.5 or 1.0
    return [value / norm for value in vector]


class HKPCEmbeddings(Embeddings):
    """OpenAI 兼容的 HKPC /embeddings 封装，供 FAISS 使用。"""

    def __call__(self, text: str) -> list[float]:
        return self.embed_query(text)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        if not texts:
            return []
        data = hkpc_post(
            "/embeddings",
            {"model": active_embedding_model(), "input": texts},
        )
        items = sorted(data.get("data") or [], key=lambda item: int(item.get("index", 0)))
        if len(items) != len(texts):
            raise RuntimeError(f"HKPC embedding 返回数量不符：期望 {len(texts)}，实际 {len(items)}")
        return [_l2_normalize(list(item.get("embedding") or [])) for item in items]

    def embed_query(self, text: str) -> list[float]:
        return self.embed_documents([text])[0]


class ProgressEmbeddings(Embeddings):
    def __init__(self, inner, batch_size: int = EMBEDDING_BATCH_SIZE):
        self.inner = inner
        self.batch_size = batch_size

    def __call__(self, text: str) -> list[float]:
        return self.embed_query(text)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        vectors: list[list[float]] = []
        total = len(texts)
        for start in range(0, total, self.batch_size):
            batch = texts[start : start + self.batch_size]
            vectors.extend(self.inner.embed_documents(batch))
            done = min(start + len(batch), total)
            STATUS.update(done=done, total=total, message=f"已向量化 {done}/{total}")
        return vectors

    def embed_query(self, text: str) -> list[float]:
        return self.inner.embed_query(text)


class HKPCReranker:
    def predict(self, pairs, show_progress_bar: bool = False):
        if not pairs:
            return []
        query = pairs[0][0]
        documents = [text[:4000] for _query, text in pairs]
        data = hkpc_post(
            "/rerank",
            {"model": active_reranker_model(), "query": query, "documents": documents},
        )
        scores = [0.0] * len(documents)
        for item in data.get("results") or []:
            index = int(item.get("index", 0))
            if 0 <= index < len(scores):
                scores[index] = float(item.get("relevance_score") or 0)
        return scores


@lru_cache(maxsize=1)
def get_embeddings():
    if use_online_vectors():
        print(f"Embedding 使用 HKPC 線上模型：{active_embedding_model()}")
        return HKPCEmbeddings()
    from langchain_huggingface import HuggingFaceEmbeddings

    device = resolve_torch_device()
    batch_size = 64 if device.startswith("cuda") else 8
    print(f"Embedding 使用本地模型 {active_embedding_model()}，设备：{describe_torch_device(device)}")
    return HuggingFaceEmbeddings(
        model_name=active_embedding_model(),
        model_kwargs={"device": device},
        encode_kwargs={"normalize_embeddings": True, "batch_size": batch_size},
    )


@lru_cache(maxsize=1)
def get_reranker():
    if use_online_vectors():
        print(f"Reranker 使用 HKPC 線上模型：{active_reranker_model()}")
        return HKPCReranker()
    import torch
    from sentence_transformers import CrossEncoder

    device = resolve_torch_device()
    print(f"Reranker 使用本地模型 {active_reranker_model()}，设备：{describe_torch_device(device)}")
    model_kwargs = {"torch_dtype": torch.float16} if device.startswith("cuda") else {}
    return CrossEncoder(
        active_reranker_model(),
        device=device,
        trust_remote_code=True,
        model_kwargs=model_kwargs,
    )


def load_knowledge_documents(library_id: str):
    from langchain_community.document_loaders import Docx2txtLoader, PyMuPDFLoader, TextLoader
    from langchain_core.documents import Document

    documents = []
    library_name = read_library_name(library_id)
    for file_path in iter_library_files(library_id):
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            loader = PyMuPDFLoader(str(file_path))
        elif suffix == ".docx":
            loader = Docx2txtLoader(str(file_path))
        elif suffix in {".xlsx", ".xlsm"}:
            documents.extend(load_spreadsheet(file_path, Document))
            continue
        elif suffix == ".pptx":
            documents.extend(load_presentation(file_path, Document))
            continue
        elif suffix in {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}:
            documents.extend(load_image(file_path, Document))
            continue
        else:
            loader = TextLoader(str(file_path), encoding="utf-8", autodetect_encoding=True)
        for document in loader.load():
            document.page_content = clean_text(document.page_content)
            if len(document.page_content) >= 40:
                documents.append(document)
    for document in documents:
        document.metadata = {
            **(document.metadata or {}),
            "library": library_id,
            "library_name": library_name,
        }
    return documents


def load_spreadsheet(file_path: Path, document_type):
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    documents = []
    for worksheet in workbook.worksheets:
        rows = [
            " | ".join(str(value).strip() for value in row if value is not None and str(value).strip())
            for row in worksheet.iter_rows(values_only=True)
        ]
        content = clean_text("\n".join(row for row in rows if row))
        if content:
            documents.append(document_type(page_content=content, metadata={
                "source": str(file_path), "sheet": worksheet.title, "kind": "table",
            }))
    return documents


def load_presentation(file_path: Path, document_type):
    from pptx import Presentation

    presentation = Presentation(file_path)
    documents = []
    for index, slide in enumerate(presentation.slides, start=1):
        content = clean_text("\n".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        ))
        if content:
            documents.append(document_type(page_content=content, metadata={
                "source": str(file_path), "slide": index, "kind": "text",
            }))
    return documents


def load_image(file_path: Path, document_type):
    from rapidocr_onnxruntime import RapidOCR

    result, _ = RapidOCR()(str(file_path))
    content = clean_text("\n".join(line[1] for line in result or []))
    if not content:
        raise RuntimeError(f"圖片中未識別到可檢索文字：{file_path.name}")
    return [document_type(page_content=content, metadata={
        "source": str(file_path), "kind": "image",
    })]


def persist_index(vectorstore, fingerprint: str, library_id: str) -> None:
    destination = library_index_dir(library_id)
    destination.mkdir(parents=True, exist_ok=True)
    native = faiss_native_dir(library_id)
    native.mkdir(parents=True, exist_ok=True)
    try:
        vectorstore.save_local(str(native))
        if not same_path(native, destination):
            copy_index_files(native, destination)
        (destination / "fingerprint.txt").write_text(fingerprint, encoding="utf-8")
        print(f"知識庫索引已保存：{destination}")
    except Exception as error:
        print(f"提示：索引已在記憶體中可用，磁碟寫入已跳過：{error}")


def load_saved_vectorstore(embeddings, library_id: str):
    from langchain_community.vectorstores import FAISS

    destination = library_index_dir(library_id)
    native = faiss_native_dir(library_id)
    if not same_path(native, destination):
        copy_index_files(destination, native)
    return FAISS.load_local(
        str(native),
        embeddings,
        allow_dangerous_deserialization=True,
    )


def build_vectorstore(library_id: str):
    from langchain_community.vectorstores import FAISS
    from langchain_core.documents import Document
    from zhihui_chunking import STATUS, chunk_content

    library_name = read_library_name(library_id)
    STATUS.update(phase="scanning", message=f"正在掃描「{library_name}」", done=0, total=len(iter_library_files(library_id)))
    documents = load_knowledge_documents(library_id)
    if not documents:
        raise RuntimeError(f"知識庫沒有可檢索正文：{library_name}")

    STATUS.update(phase="chunking", message=f"正在切分「{library_name}」", done=0, total=len(documents))
    chunks = []
    for position, document in enumerate(documents, start=1):
        metadata = document.metadata or {}
        source = str(metadata.get("source", ""))
        page = metadata.get("page")
        section = str(metadata.get("sheet") or metadata.get("slide") or "")
        kind = str(metadata.get("kind", "text"))
        for chunk in chunk_content(
            document.page_content,
            source,
            page,
            section,
            kind,
            chunk_size=RAG_CHUNK_SIZE,
            overlap=RAG_CHUNK_OVERLAP,
        ):
            if not chunk.text.strip():
                continue
            chunk_metadata = {
                **metadata,
                "source": chunk.source,
                "page": chunk.page,
                "section": chunk.section,
                "kind": chunk.kind,
                "library": library_id,
                "library_name": library_name,
            }
            chunks.append(Document(page_content=chunk.text, metadata=chunk_metadata))
        STATUS.update(
            done=position,
            current_file=source,
            message=f"「{library_name}」已切分 {len(chunks)} 個資料片段",
        )
    if not chunks:
        raise RuntimeError(f"「{library_name}」切分後沒有可用片段")
    STATUS.update(
        phase="embedding",
        done=0,
        total=len(chunks),
        message=f"正在為「{library_name}」的 {len(chunks)} 個資料片段建立向量",
    )
    vectorstore = FAISS.from_documents(
        chunks,
        ProgressEmbeddings(get_embeddings(), EMBEDDING_BATCH_SIZE),
    )
    STATUS.update(done=len(chunks), message=f"「{library_name}」已完成 {len(chunks)} 個資料片段向量化")
    return vectorstore


def get_vectorstore(library_id: str):
    global _active_library
    library_id = safe_library_id(library_id)
    fingerprint = file_fingerprint(library_id)
    library_name = read_library_name(library_id)
    if not fingerprint:
        raise RuntimeError(f"知識庫沒有支援的文檔：{library_name}")

    with _index_lock:
        generation = _generations.get(library_id, 0)
        if _stores.get(library_id) is not None and _fingerprints.get(library_id) == fingerprint:
            _statuses[library_id] = "ready"
            _errors[library_id] = ""
            return _stores[library_id]

        if index_is_current(library_id):
            try:
                _stores[library_id] = load_saved_vectorstore(get_embeddings(), library_id)
                _fingerprints[library_id] = fingerprint
                _statuses[library_id] = "ready"
                _errors[library_id] = ""
                return _stores[library_id]
            except Exception as error:
                print(f"提示：讀取「{library_name}」索引失敗，將重新建立：{error}")

        if any(status == "building" for status in _statuses.values()):
            raise RuntimeError("知識庫索引正在建立，請稍後再檢索。")

        print(f"正在建立「{library_name}」索引，快取目錄 {library_index_dir(library_id)}")
        _statuses[library_id] = "building"
        _errors[library_id] = ""
        _active_library = library_id
        STATUS.update(
            phase="scanning",
            started_at=time.monotonic(),
            done=0,
            total=len(iter_library_files(library_id)),
            current_file="",
            message=f"正在準備「{library_name}」索引",
            error="",
        )

    try:
        store = build_vectorstore(library_id)
        with _index_lock:
            if _generations.get(library_id, 0) != generation:
                print(f"提示：「{library_name}」索引已被更新的任務取代。")
                published = _stores.get(library_id)
                if published is not None:
                    return published
                raise RuntimeError("知識庫索引正在建立，請稍後再檢索。")
            _stores[library_id] = store
            _fingerprints[library_id] = fingerprint
            persist_index(store, fingerprint, library_id)
            _statuses[library_id] = "ready"
            _errors[library_id] = ""
            if _active_library == library_id:
                STATUS.update(phase="ready", message=f"「{library_name}」向量索引已就緒")
            print(f"「{library_name}」索引已就緒。")
            return store
    except Exception as error:
        with _index_lock:
            if _stores.get(library_id) is not None:
                _statuses[library_id] = "ready"
                _errors[library_id] = ""
                print(f"提示：「{library_name}」索引可先在記憶體中使用：{error}")
                return _stores[library_id]
            _statuses[library_id] = "error"
            _errors[library_id] = str(error)
            STATUS.update(phase="failed", error=str(error), message=f"「{library_name}」索引建立失敗")
        raise


def warmup_index(library_id: str | None = None) -> None:
    ids = [safe_library_id(library_id)] if library_id else libraries_with_files()
    pending = [item for item in ids if iter_library_files(item)]
    if not pending:
        return
    if use_online_vectors():
        print(f"向量服務：HKPC 線上 embedding={active_embedding_model()} rerank={active_reranker_model()}")
    else:
        print(f"向量計算裝置：{describe_torch_device(resolve_torch_device())}")
    print(f"FAISS 索引目錄：{INDEX_DIR}")
    while pending:
        current = pending.pop(0)
        try:
            get_vectorstore(current)
        except RuntimeError as error:
            if "正在建立" in str(error):
                pending.append(current)
                time.sleep(1)
                continue
            _statuses[current] = "error"
            _errors[current] = str(error)
            print(f"提示：後台建立知識庫索引失敗：{error}")
        except Exception as error:
            _statuses[current] = "error"
            _errors[current] = str(error)
            print(f"提示：後台建立知識庫索引失敗：{error}")


def ensure_loaded(library_id: str):
    library_id = safe_library_id(library_id)
    fingerprint = file_fingerprint(library_id)
    if not fingerprint:
        return None
    with _index_lock:
        if _stores.get(library_id) is not None and (
            _fingerprints.get(library_id) == fingerprint or _statuses.get(library_id) == "ready"
        ):
            _statuses[library_id] = "ready"
            return _stores[library_id]
        if not index_is_current(library_id):
            return _stores.get(library_id)
        store = load_saved_vectorstore(get_embeddings(), library_id)
        _stores[library_id] = store
        _fingerprints[library_id] = fingerprint
        _statuses[library_id] = "ready"
        _errors[library_id] = ""
        return store


def drop_library_index(library_id: str) -> None:
    global _active_library
    library_id = safe_library_id(library_id)
    with _index_lock:
        _generations[library_id] = _generations.get(library_id, 0) + 1
        _stores.pop(library_id, None)
        _fingerprints.pop(library_id, None)
        _statuses.pop(library_id, None)
        _errors.pop(library_id, None)
        folder = library_index_dir(library_id)
        if folder.exists():
            shutil.rmtree(folder, ignore_errors=True)
        if _active_library == library_id:
            _active_library = ""
            STATUS.update(
                phase="idle",
                started_at=0,
                done=0,
                total=0,
                current_file="",
                message="",
                error="",
            )


def invalidate_index(library_id: str) -> None:
    global _active_library
    library_id = safe_library_id(library_id)
    library_name = read_library_name(library_id)
    with _index_lock:
        _generations[library_id] = _generations.get(library_id, 0) + 1
        _stores.pop(library_id, None)
        _fingerprints.pop(library_id, None)
        _statuses[library_id] = "idle"
        _errors[library_id] = ""
        folder = library_index_dir(library_id)
        if folder.exists():
            shutil.rmtree(folder, ignore_errors=True)
    _active_library = library_id
    STATUS.update(
        phase="scanning",
        started_at=time.monotonic(),
        done=0,
        total=len(iter_library_files(library_id)),
        current_file="",
        message=f"正在準備「{library_name}」索引",
        error="",
    )


def retrieve_from_store(store, question: str, k: int | None = None):
    if store is None:
        return []
    k = clamp_top_k(RAG_TOP_K if k is None else k)
    try:
        pairs = store.similarity_search_with_relevance_scores(question, k=k * 3)
    except Exception:
        pairs = [(document, 0.0) for document in store.similarity_search(question, k=k * 3)]
    if not pairs:
        return []
    scores = get_reranker().predict(
        [(question, document.page_content) for document, _ in pairs],
        show_progress_bar=False,
    )
    documents = []
    for (document, _,), score in sorted(
        zip(pairs, scores), key=lambda item: float(item[1]), reverse=True
    )[:k]:
        document.metadata = {**(document.metadata or {}), "score": float(score or 0)}
        documents.append(document)
    return documents


def aggregate_index_status() -> str:
    ids = libraries_with_files()
    if not ids:
        return "idle"
    statuses = [_statuses.get(item, "idle") for item in ids]
    if any(item == "building" for item in statuses):
        return "building"
    if any(item == "error" for item in statuses) and not any(
        _stores.get(item) is not None or index_is_current(item) for item in ids
    ):
        return "error"
    if all(
        _fingerprints.get(item) == file_fingerprint(item) or index_is_current(item)
        for item in ids
    ):
        return "ready"
    return "idle"


def library_record(library_id: str) -> dict:
    files = iter_library_files(library_id)
    status = _statuses.get(library_id, "idle")
    ready = bool(files) and status != "building" and status != "error" and (
        _stores.get(library_id) is not None or index_is_current(library_id) or status == "ready"
    )
    return {
        "id": library_id,
        "name": read_library_name(library_id),
        "document_count": len(files),
        "index_status": "empty" if not files and status == "idle" else status,
        "index_ready": ready,
        "index_error": _errors.get(library_id, ""),
    }


def patch_kimi_transformers_compat() -> None:
    """Kimi-VL 按 transformers 4.50 编写，5.x 已移除 is_torch_fx_available。"""
    import transformers.utils.import_utils as import_utils

    if not hasattr(import_utils, "is_torch_fx_available"):
        import_utils.is_torch_fx_available = lambda: False


def _model_load_kwargs(torch_module, extra: dict | None = None) -> dict:
    kwargs = {
        "trust_remote_code": True,
        "device_map": "auto",
        "low_cpu_mem_usage": True,
        "attn_implementation": "eager",
    }
    if extra:
        kwargs.update(extra)
    kwargs["torch_dtype"] = torch_module.float16
    return kwargs


class LocalKimiLLM:
    def __init__(self, model, tokenizer):
        self.model = model
        self.tokenizer = tokenizer

    def invoke(self, prompt: str) -> str:
        import torch

        messages = [{"role": "user", "content": prompt}]
        text = self.tokenizer.apply_chat_template(
            messages, tokenize=False, add_generation_prompt=True
        )
        encoded = self.tokenizer(text, return_tensors="pt")
        embed = self.model.get_input_embeddings()
        weight = getattr(embed, "weight", None)
        device = weight.device if weight is not None else next(self.model.parameters()).device
        inputs = {
            key: value.to(device)
            for key, value in encoded.items()
            if key in {"input_ids", "attention_mask"} and hasattr(value, "to")
        }
        pad_token_id = self.tokenizer.pad_token_id or self.tokenizer.eos_token_id
        with torch.inference_mode():
            output_ids = self.model.generate(
                **inputs,
                max_new_tokens=500,
                do_sample=False,
                repetition_penalty=1.1,
                pad_token_id=pad_token_id,
            )
        prompt_length = inputs["input_ids"].shape[-1]
        generated = output_ids[0, prompt_length:]
        return self.tokenizer.decode(generated, skip_special_tokens=True).strip()


def warmup_local_llm_async() -> None:
    if _local_llm_status in {"loading", "ready"}:
        return
    threading.Thread(target=_safe_load_local_llm, name="zhihui-local-llm", daemon=True).start()


def _safe_load_local_llm() -> None:
    try:
        get_local_llm()
    except Exception as error:
        print(f"提示：本地模型加载失败：{error}")


def get_local_llm():
    global _local_llm, _local_llm_status, _local_llm_error
    if _local_llm is not None:
        return _local_llm
    with _local_llm_lock:
        if _local_llm is not None:
            return _local_llm
        if not path_exists(LOCAL_MODEL_PATH):
            _local_llm_status = "error"
            _local_llm_error = "LOCAL_MODEL_PATH 未配置或目錄不存在"
            raise RuntimeError(_local_llm_error)

        import torch
        from transformers import AutoModelForCausalLM, AutoTokenizer

        _local_llm_status = "loading"
        _local_llm_error = ""
        print(f"正在加载本地模型：{LOCAL_MODEL_PATH}（torch {torch.__version__} cuda={torch.cuda.is_available()}）")
        try:
            patch_kimi_transformers_compat()
            if not torch.cuda.is_available():
                raise RuntimeError(
                    "本地 Kimi-VL-A3B 需要 NVIDIA GPU 和 CUDA 版 PyTorch。"
                    f"当前为 {torch.__version__}，cuda_available=False。"
                    "请用 conda 环境 kimi_vl_env 启动：D:\\miniconda3\\envs\\kimi_vl_env\\python.exe main.py"
                )
            tokenizer = AutoTokenizer.from_pretrained(LOCAL_MODEL_PATH, trust_remote_code=True)
            model = None
            load_error = None
            try:
                from transformers import BitsAndBytesConfig

                quantization = BitsAndBytesConfig(
                    load_in_4bit=True,
                    bnb_4bit_quant_type="nf4",
                    bnb_4bit_compute_dtype=torch.float16,
                    bnb_4bit_use_double_quant=True,
                )
                model = AutoModelForCausalLM.from_pretrained(
                    LOCAL_MODEL_PATH,
                    quantization_config=quantization,
                    **_model_load_kwargs(torch),
                )
                print("本地模型已按 4bit 载入 GPU")
            except Exception as error:
                load_error = error
                print(f"提示：4bit 量化加载失败，改为 float16：{error}")
                model = AutoModelForCausalLM.from_pretrained(
                    LOCAL_MODEL_PATH,
                    **_model_load_kwargs(torch),
                )
                print("本地模型已按 float16 载入")
            if model is None:
                raise RuntimeError(str(load_error) if load_error else "本地模型加载失败")
            model.eval()
            _local_llm = LocalKimiLLM(model, tokenizer)
            _local_llm_status = "ready"
            print(f"本地模型已就绪：{describe_torch_device(resolve_torch_device())}")
            return _local_llm
        except Exception as error:
            _local_llm_status = "error"
            _local_llm_error = str(error)
            raise


def resolve_online_model(name: str | None) -> str:
    if name and name in ONLINE_MODELS:
        return name
    if not name:
        return ONLINE_MODEL_NAME if ONLINE_MODEL_NAME in ONLINE_MODELS else next(iter(ONLINE_MODELS))
    raise RuntimeError(f"不支援的線上模型：{name}")


def invoke_online_model(model_name: str, prompt: str) -> str:
    spec = ONLINE_MODELS[model_name]
    data = hkpc_post(
        "/chat/completions",
        {
            "model": model_name,
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.1,
            "max_tokens": spec["max_tokens"],
            "chat_template_kwargs": {"enable_thinking": False},
            "enable_thinking": False,
        },
        timeout=60,
        api_key=model_api_key(model_name),
        base_url=model_base_url(model_name),
    )
    message = data["choices"][0]["message"]
    return str(message.get("content") or message.get("reasoning") or "").strip()


def serialize_source(doc):
    metadata = doc.metadata or {}
    source_path = metadata.get("source", "")
    location = None
    if metadata.get("page") is not None:
        location = f"第 {metadata['page'] + 1} 頁"
    elif metadata.get("slide") is not None:
        location = f"第 {metadata['slide']} 頁投影片"
    elif metadata.get("sheet"):
        location = f"工作表：{metadata['sheet']}"
    content = doc.page_content or ""
    return {
        "type": Path(source_path).suffix.lstrip(".").upper() or "DOC",
        "title": source_title(source_path),
        "library": metadata.get("library", ""),
        "library_name": metadata.get("library_name", ""),
        "page": metadata.get("page"),
        "location": location,
        "score": metadata.get("score", 0),
        "excerpt": content[:500].strip(),
        "content": content,
    }


def render_system_prompt(context: str, question: str, custom_prompt: str | None) -> str:
    template = (custom_prompt or "").strip() or SYSTEM_PROMPT
    if "{context}" not in template or "{question}" not in template:
        template = template.rstrip() + "\n\n【參考資料】\n{context}\n\n【用戶問題】\n{question}\n"
    return template.replace("{context}", context).replace("{question}", question)


def apply_rag_settings(
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
    top_k: int | None = None,
    persist: bool = True,
) -> bool:
    """更新切片與召回數量。回傳值表示切片參數已改變，既有索引需要重建。"""
    global RAG_CHUNK_SIZE, RAG_CHUNK_OVERLAP, RAG_TOP_K
    before = chunking_token()
    size = RAG_CHUNK_SIZE if chunk_size is None else clamp_chunk_size(chunk_size)
    overlap = RAG_CHUNK_OVERLAP if chunk_overlap is None else clamp_chunk_overlap(chunk_overlap, size)
    if chunk_size is not None:
        overlap = clamp_chunk_overlap(overlap, size)
    top = RAG_TOP_K if top_k is None else clamp_top_k(top_k)
    changed = size != RAG_CHUNK_SIZE or overlap != RAG_CHUNK_OVERLAP or top != RAG_TOP_K
    RAG_CHUNK_SIZE, RAG_CHUNK_OVERLAP, RAG_TOP_K = size, overlap, top
    if persist and changed:
        os.environ["HKPC_CHUNK_SIZE"] = str(size)
        os.environ["HKPC_CHUNK_OVERLAP"] = str(overlap)
        os.environ["HKPC_TOP_K"] = str(top)
        update_env_file({
            "HKPC_CHUNK_SIZE": str(size),
            "HKPC_CHUNK_OVERLAP": str(overlap),
            "HKPC_TOP_K": str(top),
        })
    return chunking_token() != before


def apply_vector_settings(
    backend: str | None = None,
    embedding_model: str | None = None,
    reranker_model: str | None = None,
) -> bool:
    """寫入向量後端與模型。回傳值表示 embedding 或後端已變，需要重建索引。"""
    before_embed = active_embedding_model()
    before_online = use_online_vectors()
    before_rerank = active_reranker_model()
    updates: dict[str, str] = {}
    if backend and backend != os.environ.get("HKPC_VECTOR_BACKEND", "").strip():
        os.environ["HKPC_VECTOR_BACKEND"] = backend
        updates["HKPC_VECTOR_BACKEND"] = backend
    if embedding_model and embedding_model != os.environ.get("HKPC_EMBEDDING_MODEL", "").strip():
        os.environ["HKPC_EMBEDDING_MODEL"] = embedding_model
        updates["HKPC_EMBEDDING_MODEL"] = embedding_model
    if reranker_model and reranker_model != os.environ.get("HKPC_RERANKER_MODEL", "").strip():
        os.environ["HKPC_RERANKER_MODEL"] = reranker_model
        updates["HKPC_RERANKER_MODEL"] = reranker_model
    if updates:
        update_env_file(updates)
    embed_changed = active_embedding_model() != before_embed or use_online_vectors() != before_online
    rerank_changed = active_reranker_model() != before_rerank or use_online_vectors() != before_online
    if embed_changed:
        get_embeddings.cache_clear()
        resolve_torch_device.cache_clear()
    if rerank_changed:
        get_reranker.cache_clear()
    return embed_changed


def schedule_reindex() -> None:
    library_ids = libraries_with_files()
    for library_id in library_ids:
        invalidate_index(library_id)
    if library_ids:
        threading.Thread(target=warmup_index, name="zhihui-reindex", daemon=True).start()


def extract_answer(output, prompt: str) -> str:
    answer = output.content if hasattr(output, "content") else str(output)
    if not str(answer or "").strip():
        extra = getattr(output, "additional_kwargs", {}) or {}
        answer = extra.get("reasoning") or extra.get("reasoning_content") or str(output)
    text = str(answer or "").strip()
    if prompt and text.startswith(prompt):
        text = text[len(prompt):].strip()
    marker = "回答："
    if marker in text:
        text = text.split(marker, 1)[-1].strip()
    return text


def run_chat(request: ChatRequest):
    documents = []
    retrieval_skipped = False
    scope = (request.knowledge_scope or "").strip()
    active_library = ""
    top_k = clamp_top_k(RAG_TOP_K if request.top_k is None else request.top_k)
    if not request.use_knowledge or scope in NONE_SCOPES:
        context = "用戶明確選擇不使用知識庫檢索。請按一般知識以繁體中文回答，並說明本次未使用資料庫。"
    else:
        target = None if scope in ALL_SCOPES or not scope else resolve_library_id(scope)
        library_ids = [target] if target else libraries_with_files()
        active_library = target or ""
        skipped_names = []
        searched = []
        for library_id in library_ids:
            name = read_library_name(library_id)
            if not iter_library_files(library_id):
                skipped_names.append(f"{name}（尚無文檔）")
                continue
            if _statuses.get(library_id) == "error" and target:
                raise RuntimeError(f"「{name}」索引建立失敗：{_errors.get(library_id) or '未知錯誤'}")
            store = ensure_loaded(library_id)
            if store is None:
                skipped_names.append(name)
                continue
            searched.append(name)
            documents.extend(retrieve_from_store(store, request.question, top_k))
        if not target and documents:
            documents.sort(key=lambda doc: float((doc.metadata or {}).get("score") or 0), reverse=True)
            documents = documents[:top_k]
        retrieval_skipped = not searched
        if retrieval_skipped:
            waiting = "、".join(skipped_names) or "所選知識庫"
            if skipped_names and all("尚無文檔" in item for item in skipped_names):
                context = f"「{waiting}」目前沒有文檔。請用繁體中文說明該知識庫尚無文檔，不要假裝已引用原文。"
            else:
                context = f"「{waiting}」索引正在後台建立，本次回答未檢索文檔。索引完成後即可引用原文。請先以繁體中文直接回答，並說明尚未引用該知識庫。"
        else:
            body = (
                "\n\n".join(doc.page_content for doc in documents)
                if documents
                else "已啟用知識庫檢索，但沒有召回與問題直接相關的文檔片段。請明確說明該限制，不得將此情況表述為未啟用知識庫。"
            )
            if skipped_names:
                body += "\n\n【本次未檢索的知識庫】" + "、".join(skipped_names) + "。請不要把這些知識庫說成沒有資料。"
            if target:
                body = f"【目前知識庫】{read_library_name(target)}\n\n" + body
            context = body
    prompt = render_system_prompt(context, request.question, request.custom_prompt)
    chunk_changed = False
    if request.chunk_size is not None or request.chunk_overlap is not None or request.top_k is not None:
        chunk_changed = apply_rag_settings(request.chunk_size, request.chunk_overlap, top_k, persist=True)

    try:
        if request.provider == "local":
            output = get_local_llm().invoke(prompt)
            answer = extract_answer(output, prompt)
        else:
            answer = invoke_online_model(resolve_online_model(request.model), prompt)
            if not answer:
                answer = "模型沒有返回正文。已關閉思維鏈，請再試一次。"
    finally:
        if chunk_changed:
            schedule_reindex()

    return {
        "answer": answer,
        "sources": [serialize_source(doc) for doc in documents],
        "retrieval": {
            "enabled": request.use_knowledge and scope not in NONE_SCOPES,
            "skipped": retrieval_skipped,
            "source_count": len(documents),
            "index_status": aggregate_index_status(),
            "library": active_library,
            "top_k": top_k,
            "chunk_size": RAG_CHUNK_SIZE,
            "chunk_overlap": RAG_CHUNK_OVERLAP,
            "debug": bool(request.debug),
            "reindex_started": chunk_changed,
        },
    }


@app.get("/api/health")
def health():
    files = knowledge_files()
    return {
        "status": "ok",
        "models": {
            "local": {
                "name": LOCAL_MODEL_NAME,
                "ready": path_exists(LOCAL_MODEL_PATH),
                "loaded": _local_llm_status == "ready",
                "loading": _local_llm_status == "loading",
                "error": _local_llm_error,
            },
            "online": {
                "name": ONLINE_MODEL_NAME,
                "ready": bool(model_api_key(ONLINE_MODEL_NAME)),
                "options": [
                    {"id": model_id, "name": model_id, "label": spec["label"]}
                    for model_id, spec in ONLINE_MODELS.items()
                ],
            },
        },
        "knowledge_base": {
            "ready": bool(files),
            "index_ready": aggregate_index_status() == "ready",
            "index_status": aggregate_index_status(),
            "index_error": _errors.get(_active_library, ""),
            "index_directory": str(INDEX_DIR),
            "directory": str(knowledge_root()),
            "active_library": _active_library,
            "libraries": [library_record(item) for item in list_library_ids()],
            "document_count": len(files),
            "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
            "vector_models": {
                "embedding": active_embedding_model(),
                "reranker": active_reranker_model(),
                "backend": os.environ.get("HKPC_VECTOR_BACKEND", "").strip() or ("online" if use_online_vectors() else "local"),
                "device": "hkpc-online" if use_online_vectors() else resolve_torch_device(),
                "chunk_size": RAG_CHUNK_SIZE,
                "chunk_overlap": RAG_CHUNK_OVERLAP,
                "top_k": RAG_TOP_K,
            },
            "index_progress": STATUS.snapshot(),
        },
    }


@app.post("/api/local/warmup")
def warmup_local_model():
    if not path_exists(LOCAL_MODEL_PATH):
        raise HTTPException(status_code=400, detail="本地模型目錄不存在")
    warmup_local_llm_async()
    return {"status": _local_llm_status, "error": _local_llm_error}


class DevSettings(BaseModel):
    api_key: str | None = None
    base_url: str | None = None
    model: str | None = None
    vector_backend: str | None = None
    embedding_model: str | None = None
    reranker_model: str | None = None
    chunk_size: int | None = Field(default=None, ge=200, le=4000)
    chunk_overlap: int | None = Field(default=None, ge=0, le=2000)
    top_k: int | None = Field(default=None, ge=1, le=20)


class ModelWrite(BaseModel):
    id: str = Field(min_length=1, max_length=120)
    label: str = Field(min_length=1, max_length=40)
    new_id: str | None = None
    api_key: str | None = None
    base_url: str | None = None


def clean_api_key(value: str | None) -> str:
    api_key = (value or "").strip()
    if api_key and ("\n" in api_key or "\r" in api_key or len(api_key) > 500):
        raise HTTPException(status_code=400, detail="API Key 格式不正確")
    return api_key


def clean_base_url(value: str | None) -> str:
    url = (value or "").strip().rstrip("/")
    if not url:
        return ""
    if not url.startswith(("http://", "https://")) or len(url) > 300 or "\n" in url or "\r" in url:
        raise HTTPException(status_code=400, detail="URL 需要以 http:// 或 https:// 開頭")
    return url


def clean_model_id(value: str) -> str:
    model_id = (value or "").strip()
    if not MODEL_ID_RE.fullmatch(model_id):
        raise HTTPException(status_code=400, detail="模型編號格式不正確")
    return model_id


def vector_backend_label() -> str:
    stored = os.environ.get("HKPC_VECTOR_BACKEND", "").strip()
    return stored or ("online" if use_online_vectors() else "local")


def clean_vector_backend(value: str | None) -> str | None:
    text = (value or "").strip().lower()
    if not text:
        return None
    if text not in {"online", "hkpc", "cloud", "local", "cpu", "gpu"}:
        raise HTTPException(status_code=400, detail="向量後端只可填 online、hkpc、cloud、local、cpu 或 gpu")
    return text


def clean_vector_model(value: str | None, label: str) -> str | None:
    text = (value or "").strip()
    if not text:
        return None
    if "\n" in text or "\r" in text or not re.fullmatch(r"[A-Za-z0-9_./:-]{1,160}", text):
        raise HTTPException(status_code=400, detail=f"{label}格式不正確")
    return text


@app.get("/api/dev-settings")
def get_dev_settings():
    return {
        "api_key_set": bool(ONLINE_API_KEY),
        "base_url": ONLINE_BASE_URL,
        "model": ONLINE_MODEL_NAME,
        "models": public_model_rows(),
        "vector_backend": vector_backend_label(),
        "embedding_model": active_embedding_model(),
        "reranker_model": active_reranker_model(),
        "chunk_size": RAG_CHUNK_SIZE,
        "chunk_overlap": RAG_CHUNK_OVERLAP,
        "top_k": RAG_TOP_K,
    }


@app.post("/api/dev-settings")
def save_dev_settings(request: DevSettings):
    api_key = (request.api_key or "").strip()
    base_url = (request.base_url or "").strip()
    model = (request.model or "").strip()
    if api_key and ("\n" in api_key or "\r" in api_key or len(api_key) > 500):
        raise HTTPException(status_code=400, detail="API Key 格式不正確")
    if base_url and not base_url.startswith(("http://", "https://")):
        raise HTTPException(status_code=400, detail="Base URL 需要以 http:// 或 https:// 開頭")
    if model and not re.fullmatch(r"[A-Za-z0-9_./:-]{1,120}", model):
        raise HTTPException(status_code=400, detail="模型編號格式不正確")
    backend = clean_vector_backend(request.vector_backend)
    embedding_model = clean_vector_model(request.embedding_model, "Embedding 模型")
    reranker_model = clean_vector_model(request.reranker_model, "Reranker 模型")
    apply_online_settings(
        api_key or None,
        base_url or None,
        model or None,
    )
    chunk_changed = False
    if request.chunk_size is not None or request.chunk_overlap is not None or request.top_k is not None:
        chunk_changed = apply_rag_settings(request.chunk_size, request.chunk_overlap, request.top_k, persist=True)
    embed_changed = apply_vector_settings(backend, embedding_model, reranker_model)
    reindex_started = bool(chunk_changed or embed_changed)
    if reindex_started:
        schedule_reindex()
    payload = get_dev_settings()
    payload["reindex_started"] = reindex_started
    return payload


@app.post("/api/dev-models")
def add_online_model(request: ModelWrite):
    model_id = clean_model_id(request.id)
    label = request.label.strip()
    api_key = clean_api_key(request.api_key)
    base_url = clean_base_url(request.base_url)
    if model_id in ONLINE_MODELS:
        raise HTTPException(status_code=400, detail="這個模型編號已存在")
    ONLINE_MODELS[model_id] = {"label": label, "max_tokens": 512, "api_key": api_key, "base_url": base_url}
    save_model_catalog()
    return get_dev_settings()


@app.put("/api/dev-models")
def update_online_model(request: ModelWrite):
    model_id = clean_model_id(request.id)
    if model_id not in ONLINE_MODELS:
        raise HTTPException(status_code=404, detail="模型不存在")
    label = request.label.strip()
    api_key = clean_api_key(request.api_key)
    base_url = clean_base_url(request.base_url)
    new_id = clean_model_id(request.new_id) if request.new_id else model_id
    if new_id != model_id and new_id in ONLINE_MODELS:
        raise HTTPException(status_code=400, detail="這個模型編號已存在")
    previous = ONLINE_MODELS[model_id]
    commercial = bool(previous.get("commercial"))
    kept_key = api_key or str(previous.get("api_key") or "")
    rebuilt: dict[str, dict] = {}
    for key, spec in ONLINE_MODELS.items():
        if key == model_id:
            rebuilt[new_id] = {
                "label": label,
                "max_tokens": 512,
                "commercial": commercial,
                "api_key": kept_key,
                "base_url": base_url,
            }
        else:
            rebuilt[key] = spec
    ONLINE_MODELS.clear()
    ONLINE_MODELS.update(rebuilt)
    if ONLINE_MODEL_NAME == model_id and new_id != model_id:
        apply_online_settings(model=new_id)
    save_model_catalog()
    return get_dev_settings()


@app.delete("/api/dev-models")
def delete_online_model(model_id: str):
    model_id = clean_model_id(model_id)
    if model_id not in ONLINE_MODELS:
        raise HTTPException(status_code=404, detail="模型不存在")
    if len(ONLINE_MODELS) <= 1:
        raise HTTPException(status_code=400, detail="至少保留一個線上模型")
    ONLINE_MODELS.pop(model_id, None)
    save_model_catalog()
    if ONLINE_MODEL_NAME == model_id:
        apply_online_settings(model=next(iter(ONLINE_MODELS)))
    return get_dev_settings()


class LibraryCreate(BaseModel):
    name: str = Field(min_length=1, max_length=40)
    id: str | None = None


@app.get("/api/libraries")
def list_libraries():
    ensure_builtin_libraries()
    return {"libraries": [library_record(item) for item in list_library_ids()]}


@app.post("/api/libraries")
def create_library(request: LibraryCreate):
    name = request.name.strip()
    for library_id in list_library_ids():
        if read_library_name(library_id) == name:
            return library_record(library_id)
    try:
        library_id = safe_library_id(request.id) if request.id else slug_from_name(name)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    write_library_meta(library_id, name)
    return library_record(library_id)


@app.delete("/api/libraries/{library_id}")
def delete_library(library_id: str):
    try:
        library_id = safe_library_id(library_id)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    folder = library_dir(library_id)
    if not folder.is_dir():
        raise HTTPException(status_code=404, detail="知識庫不存在")
    root = knowledge_root().resolve()
    target = folder.resolve()
    if root != target and root not in target.parents:
        raise HTTPException(status_code=400, detail="知識庫路徑不正確")
    name = read_library_name(library_id)
    drop_library_index(library_id)
    shutil.rmtree(target)
    return {"id": library_id, "name": name, "status": "deleted"}


@app.get("/api/documents")
def list_documents():
    documents = []
    for library_id in list_library_ids():
        root = library_dir(library_id)
        for path in iter_library_files(library_id):
            documents.append({
                "name": path.name,
                "path": str(path.relative_to(root)),
                "library": library_id,
                "library_name": read_library_name(library_id),
                "type": path.suffix.lstrip(".").upper(),
                "size": path.stat().st_size,
            })
    return {"documents": documents, "libraries": [library_record(item) for item in list_library_ids()]}


@app.post("/api/documents")
async def upload_document(
    file: UploadFile = File(...),
    library: str = Form("moldpdf"),
    chunk_size: int | None = Form(None),
    chunk_overlap: int | None = Form(None),
):
    original_name = Path(file.filename or "").name
    suffix = Path(original_name).suffix.lower()
    if not original_name or suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="僅支援 PDF、Word（DOCX）、Excel（XLSX/XLSM）、PowerPoint（PPTX）、圖片、TXT 和 Markdown",
        )
    try:
        library_id = resolve_library_id(library)
    except RuntimeError as error:
        raise HTTPException(status_code=404, detail=str(error)) from error
    if chunk_size is not None and not 200 <= int(chunk_size) <= 4000:
        raise HTTPException(status_code=400, detail="chunk_size 需介於 200 到 4000")
    if chunk_overlap is not None and not 0 <= int(chunk_overlap) <= 2000:
        raise HTTPException(status_code=400, detail="chunk_overlap 需介於 0 到 2000")
    chunk_changed = False
    if chunk_size is not None or chunk_overlap is not None:
        chunk_changed = apply_rag_settings(chunk_size, chunk_overlap, persist=True)

    folder = library_dir(library_id)
    folder.mkdir(parents=True, exist_ok=True)
    destination = folder / original_name
    if destination.exists():
        stem, counter = destination.stem, 2
        while destination.exists():
            destination = folder / f"{stem}-{counter}{suffix}"
            counter += 1

    try:
        with destination.open("wb") as output:
            shutil.copyfileobj(file.file, output)
    finally:
        await file.close()
    if chunk_changed:
        schedule_reindex()
    else:
        invalidate_index(library_id)
        threading.Thread(target=warmup_index, args=(library_id,), name="zhihui-reindex", daemon=True).start()
    return {
        "name": destination.name,
        "library": library_id,
        "status": "uploaded",
        "chunk_size": RAG_CHUNK_SIZE,
        "chunk_overlap": RAG_CHUNK_OVERLAP,
        "reindex_started": True,
    }


@app.delete("/api/documents/{library_id}/{document_path:path}")
def delete_document(library_id: str, document_path: str):
    try:
        library_id = safe_library_id(library_id)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    root = library_dir(library_id).resolve()
    target = (root / document_path).resolve()
    if not target.is_file() or (root != target and root not in target.parents):
        raise HTTPException(status_code=404, detail="文檔不存在")
    if target.name.lower() in {"readme.md", "library.json"}:
        raise HTTPException(status_code=400, detail="不能刪除知識庫說明檔")
    target.unlink()
    if iter_library_files(library_id):
        invalidate_index(library_id)
        threading.Thread(target=warmup_index, args=(library_id,), name="zhihui-reindex", daemon=True).start()
    else:
        invalidate_index(library_id)
        _statuses[library_id] = "idle"
    return {"name": document_path, "library": library_id, "status": "deleted"}


@app.post("/api/chat")
async def chat(request: ChatRequest):
    try:
        return await asyncio.to_thread(run_chat, request)
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"模型調用失敗：{error}") from error


@app.get("/")
def index():
    return FileResponse(APP_DIR / "index.html")


@app.get("/favicon.ico")
def favicon():
    icon = APP_DIR / "favicon.svg"
    return FileResponse(icon, media_type="image/svg+xml")


app.mount("/", StaticFiles(directory=APP_DIR), name="static")
